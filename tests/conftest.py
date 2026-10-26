"""Shared fixtures. All resume content is synthetic (fictional 'Jordan Q. Sample')."""

from __future__ import annotations

import os

# keep the suite offline even if a key happens to be in the environment
os.environ["LLM_JUDGE"] = "off"

import io  # noqa: E402
from pathlib import Path  # noqa: E402

import pytest

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="session")
def resume_text() -> str:
    return (FIXTURES / "sample_resume.txt").read_text()


@pytest.fixture(scope="session")
def docx_bytes(resume_text: str) -> bytes:
    import docx

    doc = docx.Document()
    for line in resume_text.splitlines():
        if line.startswith("• "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.isupper() and line.strip():
            doc.add_heading(line.title(), level=1)
        else:
            doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="session")
def pdf_bytes(resume_text: str) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    y = 750
    for line in resume_text.splitlines():
        # Helvetica lacks the bullet glyph in some builds; use '-' bullets in the PDF.
        c.drawString(50, y, line.replace("•", "-"))
        y -= 14
        if y < 50:
            c.showPage()
            y = 750
    c.save()
    return buf.getvalue()
