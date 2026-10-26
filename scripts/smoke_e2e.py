"""Offline end-to-end smoke. Fake embedder, fallback tailor, mocked Greenhouse."""

from __future__ import annotations

import json
import tempfile
from pathlib import Path

import httpx
import respx
from fastapi.testclient import TestClient

from job_agent.api.app import create_app
from job_agent.config import Settings
from job_agent.embeddings.fake import FakeEmbedder
from job_agent.storage.pipeline import InMemoryPipelineStore
from job_agent.storage.repository import InMemoryResumeRepository

ROOT = Path(__file__).resolve().parents[1]
RESUME = (ROOT / "tests" / "fixtures" / "sample_resume.txt").read_text()

GH = {
    "jobs": [
        {
            "id": 101,
            "title": "Backend Engineer",
            "absolute_url": "https://boards.greenhouse.io/example/jobs/101",
            "updated_at": "2026-01-02T00:00:00Z",
            "company_name": "Example Corp",
            "location": {"name": "Austin, TX"},
            "content": (
                "<p>Example Corp (fictional) — Backend Engineer</p>"
                "<p>Must have: Python, FastAPI, PostgreSQL, Docker.</p>"
                "<p>Nice to have: Kubernetes, GCP.</p>"
            ),
        },
        {
            "id": 102,
            "title": "Account Executive",
            "absolute_url": "https://boards.greenhouse.io/example/jobs/102",
            "content": "<p>Must have: Salesforce. This one should be filtered out.</p>",
        },
    ]
}


def _docx(text: str) -> bytes:
    import io

    import docx

    document = docx.Document()
    for line in text.splitlines():
        if line.startswith("• "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line.isupper() and line.strip():
            document.add_heading(line.title(), level=1)
        else:
            document.add_paragraph(line)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def main() -> None:
    tmp = Path(tempfile.mkdtemp(prefix="applypilot-smoke-"))
    settings = Settings(
        database_url="", upload_dir=tmp / "up", run_log_path=tmp / "runs.jsonl",
        llm_judge="off", embedding_backend="fake", http_retries=1,
    )
    app = create_app(
        settings=settings,
        repository=InMemoryResumeRepository(),
        store=InMemoryPipelineStore(tmp / "runs.jsonl"),
        embedder=FakeEmbedder(768),
    )
    client = TestClient(app)
    health = client.get("/health").json()
    up = client.post("/resumes", files={"file": ("jordan_sample.docx", _docx(RESUME))})
    up.raise_for_status()
    resume_id = up.json()["resume_id"]

    with respx.mock:
        respx.get(url__regex=r"https://boards-api\.greenhouse\.io/.*").mock(
            return_value=httpx.Response(200, json=GH)
        )
        respx.get(url__regex=r"https://api\.lever\.co/.*").mock(
            return_value=httpx.Response(404, text="no such board")
        )
        respx.get(url__regex=r".*").mock(return_value=httpx.Response(404, text="no public email"))
        disc = client.post("/discovery", json={
            "greenhouse": ["example"],
            "lever": ["not-a-real-board"],
            "roles": ["backend"],
        })
        disc.raise_for_status()
        ran = client.post("/pipeline/run", json={"resume_id": resume_id})
        ran.raise_for_status()

    apps = client.get("/applications").json()["applications"]
    app_id = apps[0]["id"]
    draft = client.get(f"/applications/{app_id}/draft").json()
    approved = client.post(f"/applications/{app_id}/approve", json={"approved_by": "abinesh"})
    approved.raise_for_status()
    # applied is a user action, after they submit it themselves
    marked = client.post(f"/applications/{app_id}/applied")
    marked.raise_for_status()
    runs = client.get("/runs").json()["runs"]

    report = {
        "health": health,
        "resume": {"id": resume_id, "chunks": up.json()["chunk_count"], "sections": up.json()["sections"]},
        "discovery": disc.json(),
        "pipeline": ran.json()["results"],
        "draft": {
            "status_before_approve": draft["status"],
            "tailor_mode": draft["tailor_mode"],
            "eval_blocked": draft["eval_blocked"],
            "eval_failures": (draft.get("eval") or {}).get("failures"),
            "scores": {
                name: row.get("score")
                for name, row in ((draft.get("eval") or {}).get("scores") or {}).items()
            },
            "match_score": draft["match_score"],
            "missing_skills": draft["missing_skills"],
            "contacts": draft["contacts"],
            "resume_excerpt": "\n".join(draft["resume_text"].splitlines()[:12]),
            "email_subject": draft["email_subject"],
            "email_excerpt": "\n".join(draft["email_body"].splitlines()[:8]),
        },
        "after_approve": approved.json()["status"],
        "after_user_marks_applied": marked.json()["status"],
        "applied_at": marked.json()["applied_at"],
        "run_log_steps": [f"{r['step']} ok={r['ok']} {r['message'][:140]}" for r in runs],
    }
    print(json.dumps(report, indent=2, default=str))


if __name__ == "__main__":
    main()
