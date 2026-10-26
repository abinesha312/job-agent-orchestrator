"""Extract plain text lines from PDF (pypdf) and DOCX (python-docx) resumes."""

from __future__ import annotations

import io
from pathlib import PurePath

SUPPORTED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx"}


class UnsupportedFileType(ValueError):
    pass


class ResumeParseError(ValueError):
    pass


def file_type_for(filename: str) -> str:
    ext = PurePath(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(f"Unsupported file type {ext!r}; upload a .pdf or .docx")
    return SUPPORTED_EXTENSIONS[ext]


def extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:  # pypdf raises a variety of errors on malformed input
        raise ResumeParseError(f"Could not read PDF: {exc}") from exc
    return "\n".join(pages)


def extract_docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeParseError(f"Could not read DOCX: {exc}") from exc

    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("list"):
            text = f"• {text}"
        lines.append(text)
    # Tables are common in resume templates (e.g. skills grids).
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(lines)


def extract_text(data: bytes, filename: str) -> str:
    """Return the resume's raw text. Raises UnsupportedFileType / ResumeParseError."""
    if not data:
        raise ResumeParseError("Empty file")
    kind = file_type_for(filename)
    text = extract_pdf(data) if kind == "pdf" else extract_docx(data)
    if not text.strip():
        raise ResumeParseError(
            "No extractable text (scanned/image-only PDFs need OCR, which is not implemented)"
        )
    return text
