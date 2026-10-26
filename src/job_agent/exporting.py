from __future__ import annotations

from pathlib import Path


def write_resume_docx(text: str, path: Path) -> None:
    import docx

    path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    for raw in (text or "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        stripped = line.strip()
        # section headings we emit are a single caps word (SUMMARY, SKILLS, ...)
        if stripped.isupper() and " " not in stripped and len(stripped) <= 40:
            document.add_heading(stripped.title(), level=1)
            continue
        if stripped[:1] in "•-*–" or stripped.startswith("• "):
            body = stripped.lstrip("•-*– ").strip()
            try:
                document.add_paragraph(body, style="List Bullet")
            except KeyError:
                document.add_paragraph("• " + body)
            continue
        document.add_paragraph(stripped)
    document.save(path)
